# -*- coding: utf-8 -*-
"""pptx 한글 줄바꿈 후처리: 모든 글자 런·표 셀·endParaRPr 의 언어 태그를 ko-KR 로 바꾼다.

배경: pptxgenjs 는 모든 런을 lang="en-US" 로 쓴다. 그러면 PowerPoint 가 한글을 글자 단위로
끊어 "없었/다", "따/라가지"처럼 낱말이 줄 끝에서 쪼개진다.
실험(PowerPoint COM 렌더): eaLnBrk="0" 과 WordWrap 속성은 효과가 없었고,
lang="ko-KR"(altLang="en-US") 만 낱말 단위 줄바꿈을 만들었다. 기록은 references/pptx-decks.md.

사용:
  python deck_korean_wrap.py <in.pptx> [<out.pptx>]
  out 을 생략하면 <in>_kowrap.pptx 를 같은 폴더에 만든다.
  원본은 건드리지 않는다. 같은 이름이 있으면 _v02, _v03 … 으로 새 파일을 만든다 (덮어쓰기 금지).
슬라이드 본문, 그룹 안 도형, 표 셀, 발표자 노트를 모두 처리한다.
필요: python-pptx (pip install python-pptx)
"""
import pathlib
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

GROUP_SHAPE = 6   # MSO_SHAPE_TYPE.GROUP
_count = [0]


def unique_path(path):
    """덮어쓰기 금지: 같은 이름이 있으면 _v02, _v03 … 을 붙인다."""
    p = pathlib.Path(path)
    if not p.exists():
        return p
    n = 2
    while True:
        cand = p.with_name(f"{p.stem}_v{n:02d}{p.suffix}")
        if not cand.exists():
            return cand
        n += 1


def fix_text_frame(tf):
    for p in tf.paragraphs:
        for r in p.runs:
            rpr = r._r.get_or_add_rPr()
            rpr.set("lang", "ko-KR")
            rpr.set("altLang", "en-US")
            _count[0] += 1
        end = p._p.find(qn("a:endParaRPr"))
        if end is not None:
            end.set("lang", "ko-KR")


def _is_group(sh):
    try:
        return sh.shape_type == GROUP_SHAPE
    except NotImplementedError:          # 일부 도형 종류는 shape_type 을 구현하지 않는다
        return False


def walk(shapes):
    for sh in shapes:
        if _is_group(sh):
            walk(sh.shapes)
        if sh.has_text_frame:
            fix_text_frame(sh.text_frame)
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    fix_text_frame(cell.text_frame)


def main(argv):
    if argv and argv[0] in ("-h", "--help"):   # --help 를 파일 경로로 읽지 않는다
        print(__doc__)
        return 0
    if not argv:
        print(__doc__)
        return 2
    src = pathlib.Path(argv[0])
    dst = unique_path(argv[1] if len(argv) > 1 else src.with_name(f"{src.stem}_kowrap.pptx"))
    prs = Presentation(str(src))
    for layout in prs.slide_layouts:
        walk(layout.shapes)
    for slide in prs.slides:
        walk(slide.shapes)
        if slide.has_notes_slide:        # 발표자 노트의 한글도 같은 줄바꿈 문제가 있다
            walk(slide.notes_slide.shapes)
    prs.save(str(dst))
    print(f"lang=ko-KR 적용 런 {_count[0]}개 -> {dst}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
