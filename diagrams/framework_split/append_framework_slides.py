import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
FONT = "Calibri"
C = {k: RGBColor.from_string(v) for k, v in dict(ink="1A1A1A", g1="262626", g2="737373", g3="B3B3B3", g4="D9D9D9", g5="F2F2F2", bg="F7F7F7", w="FFFFFF").items()}

def blank():
    lay = [l for l in prs.slide_layouts if l.name == "Blank"][0]
    return prs.slides.add_slide(lay)

def text(s, x, y, w, h, t, size=11.5, color="g2", bold=False, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(t.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        r = p.add_run(); r.text = line; f = r.font; f.name = FONT; f.size = Pt(size); f.bold = bold; f.color.rgb = C[color]
    return tb

def title(s, t, sub):
    text(s, 0.4, 0.25, 12.5, 0.55, t, size=22, color="ink", bold=True)
    text(s, 0.4, 0.8, 12.5, 0.35, sub, size=11.5, color="g2")

def box(s, x, y, w, h, t, fill="w", line="g1", lw=1.5, dash=False, size=14, tcolor="ink", first_bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = 0.06
    sh.fill.solid(); sh.fill.fore_color.rgb = C[fill]
    sh.line.color.rgb = C[line]; sh.line.width = Pt(lw)
    if dash: sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right"): setattr(tf, m, Inches(0.08))
    for m in ("margin_top", "margin_bottom"): setattr(tf, m, Inches(0.04))
    lines = t.split("\n") if t else []
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = line; f = r.font; f.name = FONT; f.size = Pt(size); f.color.rgb = C[tcolor]
        f.bold = first_bold and i == 0 and len(lines) > 1
    return sh

def container(s, x, y, w, h, t):
    box(s, x, y, w, h, "", fill="bg", line="g2", lw=1.5)
    text(s, x + 0.15, y + 0.08, w - 0.3, 0.4, t, size=14, color="g1", bold=True)

def line(s, x1, y1, x2, y2, color="g1", lw=1.5, head=True, dash=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = C[color]; c.line.width = Pt(lw)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", dash)
    if head:
        e = etree.SubElement(ln, qn("a:tailEnd")); e.set("type", "triangle"); e.set("w", "med"); e.set("len", "med")
    return c

STAGES = ["사전 점검", "공격 생성", "자가 치유 루프 (구조 1·2)", "사후 검증 (구조 3)"]
def strip(s, active):
    xs = [0.5, 3.65, 6.8, 9.95]
    for i, x in enumerate(xs):
        on = i in active
        box(s, x, 1.3, 2.85, 0.42, STAGES[i], fill="g1" if on else "w", line="g1" if on else "g3", lw=1.0, size=12, tcolor="w" if on else "g2")
        if i < 3: line(s, x + 2.85, 1.51, xs[i + 1], 1.51, color="g3", lw=1.0)
    line(s, 0.5, 1.95, 12.8, 1.95, color="g4", lw=1.0, head=False)

# 14. 전체 흐름 개요
s = blank()
title(s, "전체 프레임워크 — 흐름 개요", "사전 점검 → 공격 생성 → 자가 치유 루프(구조 1·2) → 사후 검증(구조 3)")
box(s, 0.5, 2.6, 2.5, 2.4, "사전 점검\n결정론적 차단 문구 검증\nDrift 카나리아\n대화형 검토 도구", size=13)
box(s, 3.55, 2.6, 2.5, 2.4, "공격 생성\n레드팀 LLM이\n카테고리별\n시나리오 생성", size=13)
box(s, 6.6, 2.6, 3.2, 2.4, "자가 치유 루프\n구조 1: 유닛 A~D 파이프라인\n구조 2: 채점 → Meta-Rule 생성\n→ SRS 병합", fill="g4", lw=2.25, size=13)
box(s, 10.45, 2.6, 2.4, 2.4, "사후 검증 (구조 3)\n헬드아웃 검증\n적응형 재공격\n교차 모델 검증·통계 검정", size=13)
line(s, 3.0, 3.8, 3.55, 3.8); line(s, 6.05, 3.8, 6.6, 3.8); line(s, 9.8, 3.8, 10.45, 3.8)
text(s, 9.65, 3.35, 0.95, 0.35, "수렴 시", size=11, color="g1", align=PP_ALIGN.CENTER)
line(s, 9.2, 2.6, 9.2, 2.0, head=False); line(s, 9.2, 2.0, 7.2, 2.0, head=False); line(s, 7.2, 2.0, 7.2, 2.6)
text(s, 6.6, 1.5, 3.2, 0.4, "미수렴 시 반복 (라운드 상한까지)", size=12, color="g1", align=PP_ALIGN.CENTER)
text(s, 0.5, 5.5, 12.3, 0.8, "다음 슬라이드부터 각 단계를 상세히 나타낸다. 상단 띠의 검은 칸이 해당 슬라이드가 전체 흐름에서 차지하는 위치이다.", size=12, color="g2")

# 15. 준비 단계
s = blank()
title(s, "준비 단계 — 사전 점검과 공격 생성", "전체 흐름에서의 위치: 사전 점검 · 공격 생성")
strip(s, [0, 1])
container(s, 0.5, 2.2, 5.6, 4.8, "사전 점검 (루프 시작 전)")
for i, t in enumerate(["① 결정론적 차단 문구 사전 검증", "② Zero-cost Drift 카나리아", "③ 대화형 검토 도구"]):
    box(s, 0.8, 2.85 + i * 1.25, 5.0, 1.0, t, size=15, first_bold=False)
box(s, 6.9, 3.7, 3.4, 1.9, "공격 생성 (구조 3-①)\n레드팀 LLM이\n카테고리별 시나리오 생성", fill="g4", size=15)
box(s, 10.9, 4.05, 1.9, 1.2, "다음 단계\n자가 치유 루프", line="g2", dash=True, size=13)
line(s, 6.1, 4.65, 6.9, 4.65); line(s, 10.3, 4.65, 10.9, 4.65)

# 16. 구조 1
s = blank()
title(s, "구조 1 — 파이프라인 아키텍처 (유닛 A~D)", "전체 흐름에서의 위치: 자가 치유 루프 (구조 1·2)")
strip(s, [2])
text(s, 0.5, 2.1, 12.33, 0.4, "입력: ② 적대적 실행의 공격 문장   →   출력: 챗봇 응답 (③ 심판관 채점으로 전달)", size=12, color="g2", align=PP_ALIGN.CENTER)
container(s, 0.5, 2.6, 12.33, 2.95, "구조 1 — 파이프라인 아키텍처 (유닛 A~D)")
units = ["A. 입력 가드레일\n명백한 탈옥·사칭·\n인코딩 우회만\n1차 차단", "B. RAG 검색\n내부 지식베이스\n에서만 관련\n정보 조회",
         "C. LLM 추론\nSRS 페르소나 +\nMeta-Rule로\n응답 생성", "D. 출력 가드레일\n메타데이터·경쟁사\n노출 등\n최종 필터링"]
xs = [0.8, 3.85, 6.9, 9.95]
for i, (x, t) in enumerate(zip(xs, units)):
    box(s, x, 3.2, 2.75, 2.05, t, fill="g4" if i == 2 else "w", lw=2.25 if i == 2 else 1.5, size=15)
    if i < 3: line(s, x + 2.75, 4.22, xs[i + 1], 4.22)
box(s, 0.5, 5.9, 12.33, 1.1, "유닛 A는 명백한 우회만 차단한다. 교묘한 사회공학적 공격(가스라이팅, 어투 강요)은 의도적으로 통과시켜\n유닛 C의 SRS·Meta-Rule이 실제로 방어하도록 남긴다.", line="g2", dash=True, size=13, first_bold=False)
line(s, 2.17, 5.25, 2.17, 5.9, color="g2", lw=1.0, head=False, dash="sysDot")

# 17. 구조 2
s = blank()
title(s, "구조 2 — 자가 치유 폐쇄 루프", "전체 흐름에서의 위치: 자가 치유 루프 (구조 1·2)")
strip(s, [2])
container(s, 0.5, 2.15, 12.33, 4.45, "구조 2 — 자가 치유 폐쇄 루프")
box(s, 0.8, 2.65, 3.0, 1.15, "① SRS 정의\n고객사 페르소나 +\n제약 조건", size=14)
box(s, 5.15, 2.65, 3.0, 1.15, "⑤ SRS 병합\n헤더·푸터에 삽입\nSRS v(n) → v(n+1)", size=14)
box(s, 9.5, 2.65, 3.0, 1.15, "④ Meta-Rule 생성\n실패 사례를 분석해\n방어 규칙 자동 작성", fill="g4", lw=2.25, size=14)
box(s, 0.8, 5.2, 3.0, 1.15, "② 적대적 실행\n공격 시나리오로\n챗봇 응답 생성", size=14)
box(s, 5.15, 5.2, 3.0, 1.15, "구조 1 통과\n유닛 A → B → C → D", line="g2", dash=True, size=14)
box(s, 9.5, 5.2, 3.0, 1.15, "③ 심판관 채점\n무상태 LLM이\nAction Matrix로 판정", fill="g4", lw=2.25, size=14)
box(s, 4.9, 3.97, 3.5, 1.05, "Meta-Rule은 SRS 프롬프트의 헤더+푸터에 이중 삽입된다.\n심판관 채점 프롬프트에는 비용 절감을 위해 한 번만 삽입한다.", line="g2", dash=True, size=11, first_bold=False)
line(s, 6.65, 3.8, 6.65, 3.97, color="g2", lw=1.0, head=False, dash="sysDot")
line(s, 2.3, 3.8, 2.3, 5.2); line(s, 3.8, 5.775, 5.15, 5.775); line(s, 8.15, 5.775, 9.5, 5.775)
line(s, 11.0, 5.2, 11.0, 3.8); text(s, 11.1, 4.3, 1.3, 0.4, "미수렴 시", size=12, color="g1")
line(s, 9.5, 3.225, 8.15, 3.225); line(s, 5.15, 3.225, 3.8, 3.225)
text(s, 3.8, 3.3, 1.35, 0.6, "라운드 상한까지\n반복", size=10, color="g1", align=PP_ALIGN.CENTER)
line(s, 11.0, 6.35, 11.0, 6.8)
box(s, 8.8, 6.8, 4.03, 0.5, "수렴 시 (전원 FULL_DEFENSE) → 사후 검증", size=12, first_bold=False)

# 18. 구조 3
s = blank()
title(s, "구조 3 — 사후 검증", "전체 흐름에서의 위치: 사후 검증 (구조 3)")
strip(s, [3])
box(s, 0.5, 3.6, 1.7, 1.5, "수렴 시\n(전원\nFULL_DEFENSE)", fill="g4", line="g2", size=12, first_bold=False)
container(s, 2.5, 2.3, 10.33, 4.3, "구조 3 — 사후 검증 (v_final 대상)")
steps = ["구조 3-③\n헬드아웃 검증\n치유에 쓰지 않은\n별도 문항으로 재검증", "구조 3-④\n적응형 재공격\n블랙박스·화이트박스\n두 위협 모델로 재공격",
         "구조 3-⑤\n교차 모델 검증\n3개 벤더로\n응답만 교체\n(심판관은 고정)", "구조 3-⑥\n통계 검정\nWilcoxon·\nMann-Whitney·\nKruskal-Wallis·\n카이제곱"]
xs = [2.75, 5.3, 7.85, 10.4]
for i, (x, t) in enumerate(zip(xs, steps)):
    box(s, x, 3.0, 2.3, 3.2, t, size=14)
    if i < 3: line(s, x + 2.3, 4.6, xs[i + 1], 4.6)
line(s, 2.2, 4.35, 2.75, 4.35)

prs.save(OUT)
print("slides:", len(prs.slides))
