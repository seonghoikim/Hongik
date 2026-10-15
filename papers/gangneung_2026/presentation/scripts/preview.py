# -*- coding: utf-8 -*-
"""Approximate PNG preview of a python-pptx deck (LibreOffice cannot open files in this
environment). Draws solid shapes, pictures and wrapped text at their real positions and flags
text boxes whose wrapped text is taller than the box.

    python3 preview.py deck.pptx out_dir
"""
import os
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

DPI = 110
FONT_R = "/usr/share/fonts/truetype/nanum/NanumGothic.ttf"
FONT_B = "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf"


def px(emu):
    return int(emu / 914400 * DPI)


def wrap(draw, s, font, width):
    lines, cur = [], ""
    for ch in s:
        if draw.textlength(cur + ch, font=font) > width and cur:
            cut = cur.rfind(" ")
            if cut > 0:
                lines.append(cur[:cut])
                cur = cur[cut + 1:] + ch
            else:
                lines.append(cur)
                cur = ch
        else:
            cur += ch
    lines.append(cur)
    return lines


def render(path, out_dir):
    prs = Presentation(path)
    W, H = px(prs.slide_width), px(prs.slide_height)
    os.makedirs(out_dir, exist_ok=True)
    problems = []
    for n, slide in enumerate(prs.slides, 1):
        im = Image.new("RGB", (W, H), "#F8F8F8")
        d = ImageDraw.Draw(im)
        for sh in slide.shapes:
            x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                from io import BytesIO
                pic = Image.open(BytesIO(sh.image.blob)).convert("RGB").resize((max(w, 1), max(h, 1)))
                im.paste(pic, (x, y))
                if x < 0 or y < 0 or x + w > W + 1 or y + h > H + 1:
                    problems.append(f"slide {n}: picture outside slide")
                continue
            if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                try:
                    c = sh.fill.fore_color.rgb
                    d.rectangle([x, y, x + w, y + h], fill="#" + str(c))
                except Exception:
                    d.rectangle([x, y, x + w, y + h], outline="#888888")
            if getattr(sh, "has_table", False) and sh.has_table:
                tb = sh.table
                ty = y
                for row in tb.rows:
                    rh = px(row.height)
                    tx = x
                    for ci, cell in enumerate(row.cells):
                        cw = px(tb.columns[ci].width)
                        d.rectangle([tx, ty, tx + cw, ty + rh], outline="#555555",
                                    fill="#44536A" if ty == y else "#FFFFFF")
                        f = ImageFont.truetype(FONT_R, int(10 * DPI / 72))
                        lines = wrap(d, cell.text, f, cw - 6)
                        for li, line in enumerate(lines):
                            d.text((tx + 3, ty + 2 + li * 14), line, font=f, fill="#FFFFFF" if ty == y else "#111111")
                        if len(lines) * 14 > rh + 4:
                            problems.append(f"slide {n}: table cell wraps beyond row: {cell.text[:20]!r}")
                        tx += cw
                    ty += max(rh, 1)
                continue
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            ty = y + 2
            for p in sh.text_frame.paragraphs:
                runs = p.runs
                if not runs:
                    ty += 10
                    continue
                size = runs[0].font.size.pt if runs[0].font.size else 18
                bold = bool(runs[0].font.bold)
                font = ImageFont.truetype(FONT_B if bold else FONT_R, int(size * DPI / 72))
                col = "#" + str(runs[0].font.color.rgb) if runs[0].font.color and runs[0].font.color.type else "#111111"
                s = "".join(r.text for r in runs)
                lh = int(size * DPI / 72 * 1.25)
                for line in wrap(d, s, font, w - 6):
                    tw = d.textlength(line, font=font)
                    lx = x + 3 if p.alignment in (None, 1) else (x + (w - tw) / 2 if p.alignment == 2 else x + w - tw - 3)
                    d.text((lx, ty), line, font=font, fill=col)
                    ty += lh
            if ty > y + h + 4:
                problems.append(f"slide {n}: text overflows box by {ty - y - h}px: {sh.text_frame.text[:30]!r}")
            if x + w > W + 1 or ty > H:
                problems.append(f"slide {n}: text outside slide: {sh.text_frame.text[:30]!r}")
        im.save(os.path.join(out_dir, f"s{n:02d}.png"))
    # contact sheets of 6 slides
    files = sorted(f for f in os.listdir(out_dir) if f.startswith("s") and f.endswith(".png"))
    for k in range(0, len(files), 6):
        sub = [Image.open(os.path.join(out_dir, f)) for f in files[k:k + 6]]
        g = Image.new("RGB", (W * 2 + 10, (H + 10) * 3), "white")
        for i, s in enumerate(sub):
            g.paste(s, ((i % 2) * (W + 10), (i // 2) * (H + 10)))
        g.save(os.path.join(out_dir, f"sheet_{k // 6 + 1}.png"))
    print("\n".join(problems) or "no overflow detected")


if __name__ == "__main__":
    render(sys.argv[1], sys.argv[2])
