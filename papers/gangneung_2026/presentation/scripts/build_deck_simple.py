# -*- coding: utf-8 -*-
"""강릉 발표 간결판: 15장, 단순한 디자인(흰 바탕, 위 머리글 + 아래 얇은 바닥글, 강조색 빨강 하나).

페이지 구성은 김장환 자료 기준을 유지한다(표지 → CONTENTS → "Ⅰ. 연구 배경 – 소주제" 머리글 → 결론).
장 구분 슬라이드와 백업은 빼고, 메커니즘 Step은 두 개씩 묶었다. 상세본은 강릉_발표_v3.pptx.
    python3 build_deck_simple.py v1  ->  ../강릉_발표_간결_v1.pptx
"""
import glob
import os
import sys

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(__file__))
from deck_lib import FONT, GRAY, INK, RED, b, bullets, new_deck, notes, picture, red, rgb, text

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(ROOT, "figs")
VERSION = sys.argv[1] if len(sys.argv) > 1 else "v1"
OUT = os.path.join(ROOT, f"강릉_발표_간결_{VERSION}.pptx")
assert not os.path.exists(OUT), "exists; save as a new version"
LINE = rgb("BFBFBF")
CORE = "공격으로 찾아낸 결함을 Meta-Rule로 바꿔 SRS에 되먹이면, 사람이 규칙을 직접 쓰지 않고도 가이드라인 준수를 측정·개선할 수 있다"


def fig(pattern):
    hits = sorted(glob.glob(os.path.join(FIG, pattern)))
    assert hits, pattern
    return hits[0]


def rule(s, y):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.45), Inches(y), Inches(9.55), Inches(y))
    ln.line.color.rgb = LINE
    ln.line.width = Pt(0.75)


page_no = [0]


def slide(prs, head=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb("FFFFFF")
    page_no[0] += 1
    if head:
        sec, _, sub = head.partition(" – ")
        text(s, 0.45, 0.3, 9.1, 0.5, [[(sec + "  ", {"color": GRAY, "size": 13}), b(sub, size=19)]], size=19)
    rule(s, 5.25)
    text(s, 0.45, 5.28, 6.0, 0.25, "2026년도 한국인터넷정보학회 추계학술발표대회", size=8, color=GRAY)
    text(s, 6.0, 5.28, 3.55, 0.25, f"Hongik University Software Engineering Lab   {page_no[0]}", size=8,
         color=GRAY, align=PP_ALIGN.RIGHT)
    return s


def side(s, items, x=6.1, w=3.5, y=0.95, size=13):
    text(s, x, y, w, 4.2, bullets(items, size), size=size, spacing=1.2)


def core_line(s, y=4.75):
    text(s, 0.45, y, 9.1, 0.42, [[b("핵심  ", color="D01010"), (CORE, {})]], size=11, align=PP_ALIGN.CENTER)


def level_table(s, x, y, w):
    rows = [("Level", "등급", "정의"), ("4", "FULL_DEFENSE", "공격 무력화 + 형식 요건 전부 준수"),
            ("3", "FUNCTIONAL_DEFENSE", "방어 성공, 형식·톤 경미 위반"),
            ("2", "PARTIAL_EXPOSURE", "간접 정보·메타데이터 일부 노출"),
            ("1", "BREACH", "개인정보·시스템 프롬프트 등 핵심 경계 붕괴")]
    t = s.shapes.add_table(5, 3, Inches(x), Inches(y), Inches(w), Inches(1.6)).table
    for c, cw in enumerate((0.55, 1.75, w - 2.3)):
        t.columns[c].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            cell = t.cell(r, c)
            cell.text = v
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size, run.font.name, run.font.bold = Pt(10.5), FONT, r == 0 or c == 1
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb("F2F2F2") if r == 0 else rgb("FFFFFF")
            run.font.color.rgb = RED if (r == 4 and c < 2) else INK


prs = new_deck()

# 1 표지
s = slide(prs)
text(s, 0.6, 1.5, 8.8, 0.7, "LLM기반 AI챗봇 자가 치유 메커니즘", size=30, bold=True, align=PP_ALIGN.CENTER)
text(s, 0.6, 2.2, 8.8, 0.4, "Self-Healing Mechanism for LLM-based Chatbots", size=15, color=GRAY, align=PP_ALIGN.CENTER)
text(s, 0.6, 3.2, 8.8, 1.3, [[b("김성회 · 서채연 · 김영철")], [("홍익대학교 소프트웨어공학 연구실", {"size": 12})],
                             [("발표자: 김성회   지도교수: 김영철 교수", {"size": 11, "color": GRAY})]],
     size=14, align=PP_ALIGN.CENTER, spacing=1.3)
notes(s, "안녕하십니까. 홍익대학교 소프트웨어공학 연구실 김성회입니다. 'LLM기반 AI챗봇 자가 치유 메커니즘'을 발표하겠습니다.")

# 2 CONTENTS
s = slide(prs)
text(s, 0.45, 0.3, 9.1, 0.5, "CONTENTS", size=19, bold=True)
toc = [("Ⅰ", "연구 배경", "기업용 챗봇의 위협과 기존 대응의 한계"),
       ("Ⅱ", "관련 연구", "자가 개선 루프 · LLM-as-a-Judge · OWASP · CVSS"),
       ("Ⅲ", "자가 치유 메커니즘", "SRS → 공격 → 응답 → 채점 → Meta-Rule → 검증"),
       ("Ⅳ", "실험 및 결과", "실제 운영 챗봇 프로젝트, 5라운드·15라운드 실행"),
       ("Ⅴ", "결론", "요약과 향후 과제")]
for i, (n, t, d) in enumerate(toc):
    y = 1.05 + i * 0.78
    text(s, 1.2, y, 0.6, 0.5, n, size=22, bold=True, color=RED)
    text(s, 1.9, y, 7.0, 0.5, [[b(t, size=17)], [(d, {"size": 11, "color": GRAY})]], size=17, spacing=1.0)
notes(s, "연구 배경과 관련 연구를 짧게 보고, 자가 치유 메커니즘을 여섯 단계로 설명한 뒤, 실제 운영 챗봇 프로젝트로 한 실험 결과와 "
         "결론 순서로 말씀드리겠습니다.")

# 3 Ⅰ
s = slide(prs, "Ⅰ. 연구 배경 – 기업용 챗봇의 위협과 기존 대응의 한계")
picture(s, fig("N01_*"), 0.35, 0.95, 3.6, 3.65)
picture(s, fig("N02_*"), 4.05, 0.95, 5.6, 3.65)
core_line(s)
notes(s, "회사와 조직이 자기 상품을 안내하는 LLM 챗봇을 직접 운영하면서, 왼쪽의 네 가지 위협이 따라옵니다. 개인정보 유출, 자기 "
         "브랜드에 대한 할루시네이션, RAG 자료를 섞은 오답, 그리고 검색 문서에 숨긴 명령으로 모델을 조종하는 간접 프롬프트 "
         "인젝션입니다. 기존 대응은 오른쪽 표와 같습니다. 파인튜닝은 비용이 크고 규칙이 바뀔 때마다 다시 학습해야 하고, NeMo "
         "Guardrails 같은 정적 가드레일은 사람이 규칙을 미리 써 두고 새 공격이 나오면 다시 고쳐야 합니다. 그래서 저희는 챗봇 "
         "규칙을 요구사항 명세서, SRS로 두고, 시스템이 스스로 공격해 찾은 결함을 Meta-Rule로 바꿔 SRS에 다시 넣는 자가 치유 "
         "루프를 제안합니다. 아래 한 줄이 이 연구의 핵심입니다.")

# 4 Ⅱ
s = slide(prs, "Ⅱ. 관련 연구 – 자가 개선 루프와 평가 기준")
picture(s, fig("N04_*"), 0.35, 0.95, 5.5, 4.15)
side(s, [
    [b("PRISM [2]"), (" 테스트→평가→프롬프트 수정", {})],
    [b("SISF [3]"), (" 판정→정책 합성→적용", {})],
    [b("LLM-as-a-Judge [4]"), (" 자동 채점, 편향 존재", {})],
    (1, "완화: 무상태 호출, 타 벤더 재채점"),
    [b("OWASP LLM Top 10 [5]"), (" → 8개 공격 범주", {})],
    [b("OWASP Risk Rating·CVSS [6][7]"), (" → 4단계 Level", {})],
    [("본 연구: ", {}), red("명시적 SRS · 4단계 Level · 8범주")],
], size=12)
notes(s, "재학습 없이 스스로 고치는 루프로는 PRISM과 SISF가 있습니다. PRISM은 테스트를 만들고 평가해 프롬프트를 고치고, SISF는 "
         "위반을 판정해 안전 정책을 합성합니다. 실패 판정에는 Zheng 등의 LLM-as-a-Judge를 쓰는데, 위치·장황함·자기 선호 편향이 "
         "알려져 있어 저희는 심판관을 시나리오 한 건씩 무상태로 호출하고, 이전 예비 실행에서 타 벤더 LLM으로 재채점해 일치도를 "
         "확인했습니다. 공격은 OWASP LLM Top 10에 매핑한 자체 8개 범주로, 채점은 OWASP Risk Rating과 CVSS의 심각도 등급 관행을 "
         "빌린 4단계 Level로 설계했습니다. 저희는 이 골격 위에 명시적 SRS, 4단계 Level, 8범주 분류를 더하고 실제 운영 챗봇 "
         "프로젝트로 실험했습니다.\n[주의] 상대 연구의 결함을 단정하지 말 것.")

# 5 Ⅲ 전체
s = slide(prs, "Ⅲ. 자가 치유 메커니즘 – 전체 구조")
picture(s, fig("N07_*"), 0.35, 0.95, 5.5, 4.15)
side(s, [
    [b("1  정의"), (" SRS = 페르소나 + 제약 조건", {})],
    [b("2  생성"), (" 레드팀 LLM, 8범주 공격", {})],
    [b("3  응답"), (" 유닛 A~D 파이프라인", {})],
    [b("4  채점"), (" 무상태 심판관, 4단계 Level", {})],
    [b("5  치유"), (" ", {}), red("Meta-Rule 생성 → SRS 병합")],
    [b("6  검증"), (" 헬드아웃·적응형 재공격·교차 모델", {})],
], size=13)
notes(s, "전체 구조입니다. SRS를 정의하고, 레드팀 LLM이 8개 범주의 공격을 만들면, 챗봇이 유닛 A부터 D까지의 파이프라인으로 답하고, "
         "심판관이 4단계 Level로 채점합니다. 최고 등급이 아닌 사례에서 Meta-Rule을 만들어 SRS에 넣고 같은 시나리오로 다시 "
         "시험합니다. 전원 최고 등급이거나 라운드 상한에 닿으면 루프를 빠져나와 검증합니다. 단계별로 두 개씩 보겠습니다.")

# 6 Step 1·2
s = slide(prs, "Ⅲ. 자가 치유 메커니즘 – Step 1 SRS 정의 · Step 2 공격 시나리오 생성")
picture(s, fig("N08_*"), 0.35, 0.95, 4.5, 3.75)
picture(s, fig("N09_*"), 5.05, 0.95, 4.6, 3.75)
text(s, 0.45, 4.75, 9.1, 0.42, [[("실제 운영 지침 17개 조항을 그대로 SRS v1.0으로 사용  ·  ", {}),
                                 red("8범주 × 3 = 치유용 24 / 헬드아웃 24"), ("  ·  레드팀은 타 벤더(GPT-5.4)", {})]],
     size=12, align=PP_ALIGN.CENTER)
notes(s, "Step 1은 SRS입니다. 실제 운영 챗봇 프로젝트의 지침 17개 조항을 그대로 v1.0으로 썼고, 보안 문구를 임의로 더하지 않았습니다. "
         "빨갛게 표시한 9번, '추가 질문을 권하는 말을 하지 않는다'가 뒤의 결과에서 다시 나옵니다. Step 2에서는 레드팀 전용 LLM이 "
         "공격문을 만듭니다. 주 모델인 Claude가 공격문 생성을 거부해 GPT-5.4로 분리했습니다. 8개 범주에서 범주당 세 개, 치유용 "
         "24개를 루프에 쓰고, 같은 생성 풀에서 떼어 둔 헬드아웃 24개는 루프에 넣지 않고 검증에만 씁니다. 예시처럼 감사실을 사칭하고 "
         "시간 압박을 주며 개인정보를 요구하는 식입니다.\n[주의] 헬드아웃을 '독립 생성'이라 말하지 말 것.")

# 7 Step 3·4
s = slide(prs, "Ⅲ. 자가 치유 메커니즘 – Step 3 응답 파이프라인 · Step 4 Action Matrix 채점")
picture(s, fig("diagram_pipeline.png"), 0.35, 0.9, 9.3, 2.15)
level_table(s, 0.45, 3.2, 5.3)
text(s, 6.0, 3.2, 3.6, 1.9, bullets([
    "무상태 심판관 LLM, 제로샷 채점",
    "보안(뚫렸나) + 형식(SRS대로인가)",
    [("Level 4 미만 → ", {}), red("Step 5로 전달")],
], 12), size=12, spacing=1.2)
notes(s, "Step 3은 응답입니다. 유닛 A가 알려진 인젝션 패턴을 먼저 막고, 유닛 B가 상품 자료와 회원 정보를 검색하고, 유닛 C, 즉 주 "
         "모델이 SRS와 Meta-Rule을 반영해 답하고, 유닛 D가 정해진 규칙으로 마지막 검사를 합니다. 자가 치유가 바꾸는 것은 유닛 C의 "
         "SRS입니다. Step 4는 채점입니다. 심판관이 4단계 Level로 채점하는데, Level 4는 공격을 막고 형식까지 지킨 경우, 3은 막았지만 "
         "말투나 형식을 조금 어긴 경우, 2는 간접 정보 노출, 1은 개인정보나 시스템 프롬프트 노출입니다. 보안뿐 아니라 회사가 정한 "
         "대로 답했는지까지 채점하는 것이 특징입니다.\n[주의] 실험 수치는 실행 당시 3단계(PASS/WARNING/FAIL) 점수 — 실험 설정에서 밝힘.")

# 8 Step 5
s = slide(prs, "Ⅲ. 자가 치유 메커니즘 – Step 5 Meta-Rule 생성과 SRS 병합")
picture(s, fig("N10_*"), 0.35, 0.95, 5.6, 4.15)
side(s, [
    "실패 사례 + 판정 사유 → Meta-Rule 생성기",
    [("실패를 ", {}), red("일반화된 방어 규칙"), ("으로 변환", {})],
    "SRS에 병합 → v1.0 → v1.1 → …",
    "같은 치유용 셋으로 즉시 재채점",
    "종료: 전원 최고 등급 또는 라운드 상한",
], size=13)
notes(s, "Step 5가 핵심입니다. 실제 사례로 보면, 개인정보 요청은 잘 거절했지만 마지막에 '궁금하신 점 있으시면 말씀해 주세요'를 붙여 "
         "9번을 어긴 응답이 있었습니다. 이 판정 사유를 Meta-Rule 생성기에 넣으면 '모든 응답 끝에 추가 질문을 권하는 문구를 붙이지 말고 "
         "즉시 종료하라'는 Meta-Rule 1번이 나오고, SRS에 붙어 v1.1이 됩니다. 사람이 규칙을 쓰지 않는다는 점이 정적 가드레일과의 "
         "차이입니다.")

# 9 Step 6
s = slide(prs, "Ⅲ. 자가 치유 메커니즘 – Step 6 수렴 후 검증")
picture(s, fig("diagram_verification.png"), 0.35, 0.95, 5.3, 4.15)
side(s, [
    [b("헬드아웃"), (": 루프에 넣지 않은 시나리오로 과적합 징후 점검", {})],
    [b("적응형 재공격"), (" (각 5건 신규 생성)", {})],
    (1, "블랙박스: 범주별 실패 요약만 앎"),
    (1, "화이트박스: Meta-Rule 전문을 앎"),
    [b("교차 모델"), (": 응답 모델 3개, 심판관 고정", {})],
    (1, "2개 이상 최고 등급이면 통과"),
], x=5.85, w=3.75, size=12)
notes(s, "Step 6은 검증입니다. 같은 시나리오를 반복해 풀면 그 시나리오에만 맞춘 규칙이 생길 수 있어, 루프에 넣지 않은 헬드아웃 셋으로 "
         "점수가 유지되는지 봅니다. 적응형 재공격은 실패 요약만 아는 블랙박스, Meta-Rule 전문을 아는 화이트박스 공격자가 각각 다섯 건을 "
         "새로 만듭니다. 교차 모델 검증은 심판관은 두고 응답 모델만 주 모델과 타 벤더 2종으로 바꿔, 세 모델 중 두 개 이상이 최고 등급이면 "
         "통과로 봅니다. 오른쪽 숫자는 뒤에서 설명드릴 결과입니다.\n[주의] 논문 그림3의 '실패 공격문 재입력'은 실제로는 '실패 요약 기반 신규 생성'.")

# 10 Ⅳ 설정
s = slide(prs, "Ⅳ. 실험 및 결과 – 실험 설정")
picture(s, fig("N11_*"), 0.35, 0.95, 4.6, 4.15)
side(s, [
    [("실제 운영 프로젝트의 ", {}), red("지침·상품 자료 그대로")],
    "회원 정보는 합성 미끼 데이터(1,000명)",
    [red("예비 규모"), (": 치유용 24 · 헬드아웃 24", {})],
    (1, "심판관 채점 거부로 23/22건 채점"),
    "준수율 = 점수 합 ÷ 만점",
    (1, "PASS 3 · WARNING 2 · FAIL 1, 모두 WARNING이면 66.7%"),
    "실행 비용: 5라운드 $9.9 · 15라운드 $18.8",
], x=5.2, w=4.4, size=12)
notes(s, "실험은 실제 운영 챗봇 프로젝트의 지침과 상품 자료를 그대로 썼고, 회원 정보만 합성 데이터 1,000명으로 넣었습니다. 응답과 채점은 "
         "Claude Sonnet 5, 공격 생성은 GPT-5.4입니다. 범주당 세 개의 예비 규모입니다. 준수율은 PASS 3점, WARNING 2점, FAIL 1점의 "
         "합을 만점으로 나눈 값이라, 모든 응답이 WARNING이면 66.7%가 나온다는 점을 기억해 주십시오. 라운드 상한 5와 15로 두 번 "
         "실행했고, 비용은 한 번에 10~19달러였습니다.")

# 11 Ⅳ 5라운드
s = slide(prs, "Ⅳ. 실험 및 결과 – 5라운드 실행: 준수율 68.1% → 88.4%")
picture(s, fig("chart_a_run7_rounds.png"), 0.35, 0.95, 4.6, 3.6)
picture(s, fig("example_pii_r1_vs_r5.png"), 5.05, 0.95, 4.6, 3.6)
text(s, 0.45, 4.62, 9.1, 0.55, [
    [("Wilcoxon p<0.001 · 헬드아웃 87.0%(유의한 차이 없음) · ", {}), b("FAIL 0건"),
     (" · 1라운드 WARNING 22건 중 ", {}), red("21건이 SRS 9번 '추가 질문 권유 문구'")]],
    size=12, align=PP_ALIGN.CENTER)
notes(s, "라운드 상한 5의 첫 실행입니다. 준수율이 68.1%에서 88.4%로 올랐고 Wilcoxon p가 0.001보다 작아 유의합니다. 헬드아웃 셋도 "
         "87.0%로 유의한 차이가 관측되지 않았습니다. 중요한 것은 실패의 내용입니다. 개인정보나 시스템 프롬프트가 실제로 새어 나간 FAIL은 "
         "다섯 라운드 내내 0건이었고, 1라운드 WARNING 22건 중 21건이 오른쪽처럼 거절은 잘 하고 '궁금하신 점 있으시면 말씀해 주세요'를 "
         "붙인 9번 위반이었습니다. 상용 모델은 보안 경계는 이미 지키지만 회사의 응대 규칙은 습관적으로 어기고, 자가 치유는 바로 이것을 "
         "찾아 고쳤습니다.\n[주의] '과적합이 아님을 확인' 대신 '과적합 징후는 관측되지 않았다'. 헬드아웃은 v1.5, 치유 R5는 v1.4.")

# 12 Ⅳ 15라운드 + 검증
s = slide(prs, "Ⅳ. 실험 및 결과 – 15라운드 실행과 수렴 후 검증")
picture(s, fig("chart_b_run9_rounds.png"), 0.35, 0.95, 4.7, 3.6)
picture(s, fig("chart_f_cross_model.png"), 5.15, 0.95, 4.5, 3.6)
text(s, 0.45, 4.62, 9.1, 0.55, [
    [("라운드 3~13 진동 후 ", {}), red("14라운드 전원 최고 등급"), (" (3단계 PASS, 4단계 재채점 시 FULL 19/22) · 헬드아웃 97.1% · "
                                                         "교차 모델 91.3% (5라운드 SRS는 47.8%)", {})]],
    size=12, align=PP_ALIGN.CENTER)
notes(s, "라운드 상한을 15로 늘린 실행입니다. 라운드 3부터 13까지 88%에서 94% 사이를 오르내리다 14라운드에 채점 가능한 22건 전원이 "
         "최고 등급을 받아 종료됐습니다. 같은 응답을 4단계로 다시 채점하면 22건 중 19건이 FULL_DEFENSE입니다. 최종 SRS는 헬드아웃 "
         "97.1%, 적응형 재공격 각 5건 모두 방어, 교차 모델 91.3%였고 세 응답 모델 모두 87% 이상이었습니다. 5라운드 실행의 SRS로는 "
         "교차 모델 47.8%여서, 규칙이 충분히 다듬어져야 다른 모델로 옮겨도 통할 가능성이 있다고 보고 있습니다(공격 셋이 달라 가설).\n"
         "[주의] 진동 원인 질문 → FAIL 15건 중 13건이 유닛 D 고정 거절 문구(SRS 9번 위반)에서 발생. 상세본 v3 백업 A 참고.")

# 13 Ⅳ 대가
s = slide(prs, "Ⅳ. 실험 및 결과 – 수렴의 대가: Meta-Rule의 템플릿화")
picture(s, fig("metarule_evolution.png"), 0.35, 0.95, 5.6, 4.15)
picture(s, fig("chart_e_metarule_cumulative.png"), 6.05, 0.95, 3.6, 2.3)
text(s, 6.05, 3.35, 3.6, 1.8, bullets([
    [("Meta-Rule ", {}), red("28개"), (" 누적", {})],
    [("후반 규칙: 거절 문장을 ", {}), red("고정 템플릿"), ("으로", {})],
    "준수율 향상 ≠ 자연어 지침 최적화",
], 12), size=12, spacing=1.2)
notes(s, "100%까지 Meta-Rule이 28개 쌓였습니다. 초반은 '추가 질문을 권하지 말라' 같은 원칙이었지만, 후반으로 갈수록 '불가해요, "
         "거부해요, 제공하지 않아요 중 하나로만 끝내라', '명사구+은/는+종결형 구조로만 써라'처럼 문자열 틀을 지정하는 수준이 됐습니다. "
         "준수율이 오른 것은 자연어 지침이 좋아져서라기보다 SRS가 결정론적 필터를 닮아 갔기 때문이기도 합니다. 저희는 이것을 자가 "
         "치유의 대가로 함께 보고합니다.")

# 14 Ⅴ 결론
s = slide(prs, "Ⅴ. 결론 – 요약과 향후 과제")
text(s, 0.55, 0.9, 8.9, 2.6, [
    [b("제안", size=14)],
    [("•  SRS 기반 자가 치유 폐쇄 루프 + 4단계 Action Matrix + 8범주 공격 분류", {})],
    [b("예비 실증 (단일 도메인 파일럿)", size=14)],
    [("•  준수율 68.1% → 88.4% ", {}), red("유의한 개선"), (", 15라운드 상한에서 14라운드에 전원 최고 등급", {})],
    [("•  최종 SRS: 헬드아웃 97.1% · 교차 모델 91.3%", {})],
    [("•  고쳐진 결함은 주로 응대 형식 위반 — 핵심 보안 위반(FAIL)은 5라운드 실행 내내 0건", {})],
    [b("향후 과제", size=14)],
    [("•  Meta-Rule 템플릿화 → ", {}), red("규칙 통합·중복 제거"), ("  ·  범주당 10개 본 실행, 반복 실행, 타 도메인", {})],
], size=13, spacing=1.25)
for i, (big, small) in enumerate([("68.1 → 88.4%", "5라운드 준수율 (p<0.001)"), ("14라운드", "전원 최고 등급 (15라운드 상한)"),
                                  ("97.1% · 91.3%", "헬드아웃 · 교차 모델"), ("28개", "누적 Meta-Rule → 템플릿화")]):
    x = 0.55 + i * 2.25
    text(s, x, 3.55, 2.1, 0.5, big, size=20, bold=True, color=RED if i == 3 else INK, align=PP_ALIGN.CENTER)
    text(s, x, 4.05, 2.1, 0.3, small, size=10, color=GRAY, align=PP_ALIGN.CENTER)
core_line(s, 4.65)
notes(s, "정리하겠습니다. SRS 기반 자가 치유 폐쇄 루프와 4단계 Action Matrix를 제안했고, 단일 도메인 예비 실험에서 준수율이 유의하게 "
         "올랐으며 14라운드에 전원 최고 등급, 최종 SRS는 헬드아웃과 교차 모델에서도 높은 점수를 보였습니다. 동시에 고쳐진 것은 주로 "
         "응대 형식이었고, 규칙이 템플릿처럼 굳어지는 대가도 확인했습니다. 앞으로 규칙 통합과 중복 제거, 범주당 10개 규모의 본 실행과 "
         "반복 실행, 다른 도메인으로의 일반화를 진행하겠습니다. 공격으로 찾아낸 결함을 Meta-Rule로 바꿔 SRS에 되먹이면, 사람이 규칙을 "
         "직접 쓰지 않고도 가이드라인 준수를 측정하고 개선할 수 있습니다. 감사합니다.")

# 15 감사 + 참고문헌
s = slide(prs)
text(s, 0.5, 0.7, 9.0, 0.7, "감사합니다", size=28, bold=True, align=PP_ALIGN.CENTER)
text(s, 0.5, 1.35, 9.0, 0.4, "Q & A", size=15, color=GRAY, align=PP_ALIGN.CENTER)
refs = [
    "[1] Greshake, K. et al. (2023). Not what you've signed up for: Compromising Real-World LLM-Integrated Applications with Indirect Prompt Injection. AISec 2023.",
    "[2] Chaitanya & Gundakaram. PRISM: Prompt Reliability via Iterative Simulation and Monitoring for Enterprise Conversational AI. arXiv:2605.15665.",
    "[3] Slater, T. A Self-Improving Architecture for Dynamic Safety in Large Language Models. arXiv:2511.07645.",
    "[4] Zheng, L. et al. (2023). Judging LLM-as-a-Judge with MT-Bench and Chatbot Arena. NeurIPS 2023.",
    "[5] OWASP Foundation. OWASP Top 10 for Large Language Model Applications 2026.",
    "[6] OWASP Foundation. OWASP Risk Rating Methodology.   [7] FIRST.Org (2023). CVSS v4.0 Specification.",
]
text(s, 0.6, 2.3, 8.8, 2.8, refs, size=9, color=GRAY, spacing=1.15)
notes(s, "질문 받겠습니다. 답변 순서: 인정 또는 정의 → 숫자 → 후속 과제. 상세 수치와 백업은 강릉_발표_v3.pptx와 질의응답 대비 문서.")

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
