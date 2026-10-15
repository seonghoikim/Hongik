# -*- coding: utf-8 -*-
"""Slide helpers that mirror the 김장환 APIC-IST 2025 deck layout (16:9, 10 x 5.625 in):
five pastel squares at the top right, header line at the top left, dark footer bar with the
conference name (left) and the lab name (right), centered section dividers.
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt
from PIL import Image

FONT = "맑은 고딕"
W, H = 10.0, 5.625
FOOTER_TOP = 5.22
SQUARES = ["CEDFF3", "C0E3DA", "FDE892", "F9C7BE", "C4DFB1"]
DARK = RGBColor(0x44, 0x53, 0x6A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
RED = RGBColor(0xD0, 0x10, 0x10)
GRAY = RGBColor(0x59, 0x59, 0x59)
CONF_L = ("2026년도 한국인터넷정보학회", "추계학술발표대회")
CONF_R = ("Hongik University", "Software Engineering Laboratory")


def rgb(h):
    return RGBColor.from_string(h)


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    return prs


def _rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color) if isinstance(color, str) else color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, runs, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, {opts}) pieces."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        pieces = [(para, {})] if isinstance(para, str) else para
        for t, o in pieces:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.underline = o.get("underline", False)
            c = o.get("color", color)
            f.color.rgb = rgb(c) if isinstance(c, str) else c
            rpr = r._r.get_or_add_rPr()
            for tag in ("a:ea", "a:cs"):
                el = rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag[2:])
                if el is None:
                    el = rpr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag[2:], {})
                    rpr.append(el)
                el.set("typeface", FONT)
    return tb


def chrome(slide):
    """Background tint, top squares and footer bar shared by every slide."""
    bg = slide.background.fill
    bg.gradient()
    bg.gradient_angle = 90
    stops = bg.gradient_stops
    stops[0].color.rgb, stops[0].position = rgb("F3F3F3"), 0.0
    stops[1].color.rgb, stops[1].position = rgb("FFFFFF"), 1.0
    x = 6.42
    for c in SQUARES:
        _rect(slide, x, 0, 0.58, 0.19, c)
        x += 0.715
    _rect(slide, 0, FOOTER_TOP, W, H - FOOTER_TOP, DARK)
    text(slide, 0.05, FOOTER_TOP + 0.03, 5.0, 0.36, [CONF_L[0], CONF_L[1]], size=8, bold=True,
         color=rgb("FFFFFF"), spacing=1.0)
    text(slide, 5.0, FOOTER_TOP + 0.03, 4.93, 0.36, [CONF_R[0], CONF_R[1]], size=8, bold=True,
         color=rgb("FFFFFF"), align=PP_ALIGN.RIGHT, spacing=1.0)


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s)
    return s


def header(slide, t):
    text(slide, 0.45, 0.28, 9.2, 0.45, t, size=17, color=INK)


def picture(slide, path, x, y, w, h):
    """Fit an image inside the box, centered, keeping its aspect ratio."""
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                                    Inches(pw), Inches(ph))


def caption(slide, t, x, y, w):
    text(slide, x, y, w, 0.3, t, size=10, color=GRAY, align=PP_ALIGN.CENTER)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def divider(prs, label):
    s = blank(prs)
    text(s, 0.5, 2.25, 9.0, 0.8, label, size=26, bold=True, align=PP_ALIGN.CENTER)
    return s


def bullets(items, size=14):
    """items: list of str or (level, str) or (level, [(text, opts)...]). Returns paragraph runs."""
    out = []
    for it in items:
        lvl, body = (0, it) if not isinstance(it, tuple) else it
        mark = "• " if lvl == 0 else "→ "
        pieces = [(body, {})] if isinstance(body, str) else list(body)
        indent = "" if lvl == 0 else "   "
        out.append([(indent + mark, {"size": size if lvl == 0 else size - 1})] +
                   [(t, {**o, "size": o.get("size", size if lvl == 0 else size - 1)}) for t, o in pieces])
    return out


def red(t, **o):
    return (t, {"color": RED, "bold": True, **o})


def b(t, **o):
    return (t, {"bold": True, **o})
