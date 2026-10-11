# -*- coding: utf-8 -*-
"""Make thesis_summary_visual_v5.docx from v4 (v4 is read only).

Program review 2026-10-11 corrections:
  - 7차/8차 held-out, adaptive and cross-model checks ran on v1.5 (one version after the last healing round)
  - cross-model tests are repeated measures: Friedman (scores) and Cochran's Q (PASS), recomputed from the
    result files by scripts/recompute_cross_model_stats.py
  - Figure 10 is replaced by diagrams/summary_v2/plot_cross_v5.png
"""
import sys
import os
import docx
from docx.oxml.ns import qn
from PIL import Image

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = "/home/user/hongik"
SRC = os.path.join(ROOT, "thesis_summary_visual_v4.docx")
DST = os.path.join(ROOT, "thesis_summary_visual_v5.docx")
FIG10 = os.path.join(ROOT, "diagrams", "summary_v2", "plot_cross_v5.png")
assert not os.path.exists(DST), "v5 already exists; do not overwrite"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

d = docx.Document(SRC)
body = d.element.body


def txt(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def paras():
    return [e for e in body.iter(qn("w:p"))]


def one(pred):
    hits = [p for p in paras() if pred(p)]
    assert len(hits) == 1, (len(hits), [txt(h)[:30] for h in hits])
    return hits[0]


def replace_in(p, old, new):
    ts = list(p.iter(qn("w:t")))
    full = "".join(t.text or "" for t in ts)
    pos = full.find(old)
    assert pos >= 0, old[:30]
    end, cur, placed = pos + len(old), 0, False
    for t in ts:
        s, e = cur, cur + len(t.text or "")
        cur = e
        if e <= pos or s >= end:
            continue
        a, b = max(pos - s, 0), min(end - s, len(t.text))
        t.text = t.text[:a] + ("" if placed else new) + t.text[b:]
        t.set(XML_SPACE, "preserve")
        placed = True


def edit(old, new):
    replace_in(one(lambda p: old in txt(p)), old, new)


# 1) 4.1 statistical tests
edit("3개 백엔드를 비교하는 교차 모델 검증에는 Kruskal-Wallis 검정과 카이제곱 검정의 3그룹 확장을 사용한다.",
     "교차 모델 검증은 같은 헬드아웃 문항에 3개 백엔드가 각각 응답하는 반복측정 자료이므로, 점수에는 Friedman 검정을, "
     "PASS 여부에는 Cochran's Q 검정을 사용한다.")

# 2) 4.2 v_final caveat for 7차/8차
edit("만점에는 도달하지 못했다.",
     "만점에는 도달하지 못했다. 또한 상한에 도달한 뒤에도 마지막 라운드의 실패 사례로 Meta-Rule이 한 번 더 생성되어, "
     "헬드아웃·적응형 재공격·교차 모델 검증은 치유용 셋에서 평가되지 않은 v1.5로 수행되었다(8차도 같다). 따라서 7·8차의 "
     "치유용 최종 라운드(v1.4)와 헬드아웃(v1.5)의 비교는 명세서가 한 단계 다른 두 버전의 비교이다.")

# 3) 4.3 text
p43 = one(lambda p: txt(p).startswith("7~9차 세 실행 모두에서 헬드아웃 셋을 3개 백엔드에 재실행하였다"))
replace_in(p43, txt(p43),
           "7~9차 세 실행 모두에서 헬드아웃 셋을 3개 백엔드에 재실행하였다(표 7). 같은 문항에 세 백엔드가 각각 응답한 "
           "반복측정 자료이므로 점수는 Friedman 검정, PASS 여부는 Cochran's Q 검정으로 비교하였다. 8차에서는 백엔드 간 격차가 "
           "36.4%p(68.2%에서 31.8%)에 달해 점수 기준으로는 유의하였으나(Friedman p=0.036), PASS 여부 기준으로는 유의 수준에 "
           "이르지 못했다(Cochran's Q p=0.059). 7차(p=0.687, 0.565)와 9차(p=0.097, 0.097)는 두 검정 모두 유의하지 않았다. "
           "표본이 시나리오 22~23개뿐이므로 퍼센트 격차의 크기와 통계적 유의성을 혼동해서는 안 되며, 세 번의 실행 중 벤더 간 "
           "차이로 볼 근거가 있는 것은 8차 한 번뿐이고 그마저 검정 기준에 따라 경계 수준이다.")

# 4) Table 7: SRS version of 7차 and the p column
tb = next(t for t in d.tables if t.rows[0].cells[-1].text == "Kruskal-Wallis p")
NEW = {"Kruskal-Wallis p": "p (Friedman / Cochran's Q)", "v1.4 (5라운드, 원칙 수준)": "v1.5 (5라운드, 원칙 수준)",
       "0.829": "0.687 / 0.565", "0.031": "0.036 / 0.059", "0.230": "0.097 / 0.097"}
for row in tb.rows:
    for cell in row.cells:
        if cell.text in NEW:
            replace_in(cell.paragraphs[0]._p, cell.text, NEW[cell.text])

# 5) 4.4 text
edit("Kruskal-Wallis p=1.95×10⁻⁷", "Friedman·Cochran's Q 모두 p=2.98×10⁻⁶")

# 6) 5.1 text
edit("헬드아웃 셋(7차 87.0%, 9차 97.1%)의 성능이",
     "헬드아웃 셋(7차 87.0%, 9차 97.1%)의 성능이 명세서 버전 차이(7차는 헬드아웃이 한 단계 뒤인 v1.5)를 감안하더라도")

# 7) Figure 10 image
cap = one(lambda p: txt(p).startswith("그림 10. "))
pic_p = cap.getprevious()
while not list(pic_p.iter(qn("a:blip"))):
    pic_p = pic_p.getprevious()
blip = next(pic_p.iter(qn("a:blip")))
part = d.part.related_parts[blip.get(qn("r:embed"))]
with open(FIG10, "rb") as f:
    part._blob = f.read()
w, h = Image.open(FIG10).size
for ext in list(pic_p.iter(qn("wp:extent"))) + list(pic_p.iter(qn("a:ext"))):
    if ext.get("cx"):
        ext.set("cy", str(round(int(ext.get("cx")) * h / w)))

left = [k for k in ("Kruskal", "0.829", "1.95×10") if k in "".join(txt(p) for p in paras())]
assert not left, left
d.save(DST)
print("saved", DST)
