# -*- coding: utf-8 -*-
"""Research overview form (연구 개요), following the department's one-page sample:
title line, purpose and background (2 paragraphs), scope and content (first/second/third).
Output: docs/research_overview_v1.docx (never overwritten).
"""
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

DST = "/home/user/hongik/docs/research_overview_v1.docx"
assert not os.path.exists(DST), "exists; save as a new version"

TITLE = "AI 챗봇의 정보 불일치성 개선: 응대 정책 명세 기반 강건화 메커니즘 - 가전 유통 사례를 중심으로"
PURPOSE = [
    "기업이 RAG 기반 LLM 챗봇을 고객 상담에 도입하면서, 챗봇이 고객 개인정보를 노출하거나 제품 정보와 다른 답변을 하고 "
    "사용자가 주입한 허위 전제에 동조하는 정보 불일치 문제가 반복되고 있다. 지금까지는 개발자가 시스템 프롬프트를 수작업으로 "
    "작성해 일부 사례로 시험하거나 가드레일 규칙을 사람이 유지해 왔으나, 새로운 공격이 나올 때마다 규칙을 다시 써야 한다. 최근 "
    "실패 사례로 프롬프트를 자동 보강하는 연구가 제안되었지만, 개선을 하나의 통과율로 보고하여 그 개선이 정보의 진위와 기밀 "
    "측면에서 일어났는지 응답 형식에서 일어났는지 구분하지 못한다.",
    "따라서 본 연구는 고객사의 요구사항을 자연어 응대 정책으로 명세하고, 적대적 테스트로 찾은 결함을 보강 규칙으로 자동 "
    "반영하는 강건화 메커니즘을 제안한다. 정보 무질서 이론의 오정보·허위조작정보·악의정보 3분류를 조직 챗봇의 1:1 상담에 맞게 "
    "정보 불일치성으로 재정의하고, 응답을 정보 무결성과 응대 정책 준수의 두 축으로 나누어 채점하며 정답이 있는 항목은 코드로 "
    "대조한다. 이를 통해 재학습 없이 조직 맞춤형 챗봇의 정보 불일치성을 개선하고, 그 개선을 정보 무결성 기준으로 확인한다.",
]
SCOPE_LEAD = "본 연구의 범위는 크게 세 가지로 구분된다."
SCOPE = [
    "첫째, 입력 가드레일, RAG 검색, LLM 추론, 출력 가드레일의 4단계 유닛 구조 위에서 응대 정책을 자동 보강하는 강건화 루프를 "
    "설계하고 구현한다. 레드팀 LLM이 생성한 공격 문항과 제품 설명서의 사실로 만든 허위 전제·RAG 오염 문항을 OWASP LLM 10대 "
    "위협과 정보 불일치성 3분류에 대응하는 9개 범주로 구성하고, 실패 사례에서 생성한 보강 규칙을 응대 정책에 병합하는 과정을 "
    "라운드 단위로 반복한다.",
    "둘째, 무상태 LLM 심판관이 응답을 정보 무결성(유지·부분 훼손·훼손)과 응대 정책 준수(준수·경미 위반·중대 위반)의 두 축으로 "
    "채점하는 측정 체계를 구축한다. 정답이 있는 항목은 개인정보, 응대 정책 원문, 허위 값을 코드로 대조해 판정하고 심판관 판정과의 "
    "일치도를 함께 보고한다. 이 체계로 실제 운영 중인 가전 유통 챗봇의 예비 실험을 재분석하여, 단일 지표 준수율의 향상(68.2%에서 "
    "100%)이 대부분 형식 교정이었고 초기 응대 정책에서도 22건 중 21건이 정보 측면에서 이미 방어되었음을 확인한다.",
    "셋째, 보강용 셋과 분리된 헬드아웃 셋, 블랙박스·화이트박스 적응형 재공격, 응답 생성 모델만 교체하는 교차 모델 검증으로 강건화 "
    "효과의 일반화를 평가한다. 셋당 38개 문항, 5라운드, 문항당 5회 반복의 본 실험에서 정보 무결성의 개선, 헬드아웃 셋의 "
    "비열등성, 응답 생성 모델에 따른 차이를 통계적으로 검정하고, 보강 규칙 누적에 따른 문자열 템플릿화와 재현성 위협 등 "
    "부작용을 함께 분석한다.",
]
FONT = "맑은 고딕"


def set_font(run, size, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(k), FONT)


def para(cell_or_doc, text, size=10, bold=False, indent=True, align=None, space_after=0):
    p = cell_or_doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(space_after), 1.35
    if indent:
        pf.first_line_indent = Pt(size)
    if align is not None:
        p.alignment = align
    set_font(p.add_run(text), size, bold)
    return p


def borders(table):
    tbl_pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "8")
        e.set(qn("w:color"), "000000")
        b.append(e)
    tbl_pr.append(b)


doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.top_margin = sec.bottom_margin = Cm(2.0)
sec.left_margin = sec.right_margin = Cm(2.2)
doc.paragraphs[0]._p.getparent().remove(doc.paragraphs[0]._p) if doc.paragraphs else None

para(doc, "연   구   개   요", size=20, bold=True, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
table = doc.add_table(rows=3, cols=1)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
borders(table)
c1, c2, c3 = (table.rows[i].cells[0] for i in range(3))
for c in (c1, c2, c3):
    c._tc.remove(c.paragraphs[0]._p)
para(c1, f"1. 연 구 제 목 : {TITLE}", indent=False)
para(c2, "2. 연구의 목적 및 배경 :", indent=False)
for t in PURPOSE:
    para(c2, t)
para(c2, "", indent=False)
para(c3, "3. 연구의 범위 및 내용", indent=False)
para(c3, SCOPE_LEAD)
for t in SCOPE:
    para(c3, t)
para(c3, "", indent=False)
para(doc, "※ 지면이 부족할 경우 별지 이용바람", size=9, indent=False)
doc.save(DST)

chars = len(TITLE) + sum(map(len, PURPOSE)) + len(SCOPE_LEAD) + sum(map(len, SCOPE))
print("saved", DST, "| 본문 글자 수(공백 포함):", chars)
